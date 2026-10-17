#!/usr/bin/env python3
"""
Turn an official emission-factor publication into a BROC factor pack.

    # 1. see what is actually in the file
    ./build.py inspect ~/Downloads/ghg-conversion-factors-2026.xlsx

    # 2. write a mapping describing the columns you found
    # 3. emit the CSV
    ./build.py build defra-2026/mapping.json

    # 4. always
    ./verify.py defra-2026/factors.csv

Why a mapping file rather than a parser per source: publishers rearrange their
spreadsheets every year. A hard-coded parser silently reads the wrong column the
first time a sheet moves, and a wrong emission factor that parses cleanly is the
worst failure this project has. A mapping is re-checked by a human each year,
lives in the pack next to the data, and is diffable against last year's.

Requires openpyxl for .xlsx:  pip install openpyxl
"""
from __future__ import annotations
import argparse, csv, json, sys
from pathlib import Path

CATEGORY_UNIT = {'scope1': 'MMBtu', 'scope2_location': 'kWh', 'refrigerant': 'kg'}

COLUMNS = ['factor_category', 'factor_key', 'label', 'co2e_factor', 'unit',
           'factor_source', 'data_year', 'effective_from', 'effective_to', 'notes']

# Pure unit arithmetic — definitions, not emission data. A factor expressed per
# GJ or per therm is the same factor; only the denominator changes.
TO_MMBTU = {
    'mmbtu': 1.0,
    'kwh': 0.003412142,        # 1 kWh = 3412.142 Btu
    'mwh': 3.412142,
    'gj': 0.9478171,
    'mj': 0.0009478171,
    'therm': 0.1,
    'mmbtu_hhv': 1.0,
}
MASS_TO_KG = {'kg': 1.0, 'g': 0.001, 'tonne': 1000.0, 't': 1000.0, 'lb': 0.45359237, 'short_ton': 907.18474}


def load_sheet(path: Path, sheet: str | None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is not installed. pip install openpyxl")
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet is None:
        return wb, None
    if sheet not in wb.sheetnames:
        sys.exit(f"no sheet named '{sheet}'. Available: {', '.join(wb.sheetnames)}")
    return wb, wb[sheet]


def cmd_inspect(args) -> int:
    path = Path(args.path).expanduser()
    wb, _ = load_sheet(path, None)
    print(f"{path.name}\n{len(wb.sheetnames)} sheet(s)\n")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"── {name} ──")
        for r, row in enumerate(ws.iter_rows(max_row=args.rows, values_only=True), start=1):
            cells = ['' if c is None else str(c)[:26] for c in row[:args.cols]]
            if any(cells):
                print(f"  r{r}: " + ' | '.join(cells))
        print()
    print("Write a mapping.json describing the sheet, the header row, and which\n"
          "columns hold the key, the label and the factor. See README.md.")
    return 0


def convert(value: float, category: str, source_unit: str, energy_content: float | None) -> float:
    """Convert a published factor to the category's denominator."""
    su = (source_unit or '').strip().lower().replace(' ', '_')

    if category == 'refrigerant':
        # A GWP is dimensionless (kg CO2e per kg). Only a mass basis needs work.
        return value / MASS_TO_KG[su] if su in MASS_TO_KG and su != 'kg' else value

    if category == 'scope2_location':
        if su in ('kwh', ''):
            return value
        if su == 'mwh':
            return value / 1000.0
        if su in MASS_TO_KG and su != 'kg':
            # e.g. published as lb CO2e/kWh
            return value * MASS_TO_KG[su]
        raise ValueError(f"cannot convert a grid factor from '{source_unit}' to kWh")

    # scope1: target MMBtu
    if su in TO_MMBTU:
        # published per <su>; kg CO2e per MMBtu = value / (MMBtu per <su>)
        return value / TO_MMBTU[su]
    if energy_content:
        # Published per physical unit (litre, kg, m3). The energy content comes
        # from the publication's own table and must be stated in the mapping —
        # it is source data, not something this script may assume.
        return value / energy_content
    raise ValueError(
        f"'{source_unit}' is a physical unit, so the conversion needs the fuel's energy "
        f"content. Add \"energy_content_mmbtu_per_unit\" to this row group, taken from "
        f"the publication's own energy-content table."
    )


def cmd_build(args) -> int:
    mpath = Path(args.mapping).expanduser()
    m = json.loads(mpath.read_text(encoding='utf-8'))
    base = mpath.parent

    src_file = (base / m['source_file']).expanduser() if not Path(m['source_file']).is_absolute() \
        else Path(m['source_file']).expanduser()
    if not src_file.exists():
        sys.exit(f"source file not found: {src_file}\nDownload it from {m.get('source_url', '(no url in mapping)')}")

    factor_source = m['factor_source']
    data_year = int(m['data_year'])
    eff_from = m['effective_from']
    eff_to = m.get('effective_to') or ''

    out_rows: list[dict] = []
    skipped = 0

    for group in m['groups']:
        category = group['category']
        if category not in CATEGORY_UNIT:
            sys.exit(f"unknown category '{category}'")
        _, ws = load_sheet(src_file, group['sheet'])
        header_row = int(group.get('header_row', 1))
        key_col = group['key_column']
        val_col = group['value_column']
        label_col = group.get('label_column')
        source_unit = group.get('source_unit', CATEGORY_UNIT[category])
        energy_content = group.get('energy_content_mmbtu_per_unit')
        note = group.get('note') or ''
        group_source = group.get('factor_source') or factor_source

        rows = list(ws.iter_rows(min_row=header_row, values_only=True))
        if not rows:
            sys.exit(f"sheet '{group['sheet']}' has no rows at or after row {header_row}")
        header = ['' if c is None else str(c).strip() for c in rows[0]]

        def idx(name: str) -> int:
            if name in header:
                return header.index(name)
            sys.exit(f"sheet '{group['sheet']}' has no column '{name}'. Found: {', '.join(h for h in header if h)}")

        ki, vi = idx(key_col), idx(val_col)
        li = idx(label_col) if label_col else None

        for raw in rows[1:]:
            key = raw[ki] if ki < len(raw) else None
            val = raw[vi] if vi < len(raw) else None
            if key is None or val is None or str(key).strip() == '':
                skipped += 1
                continue
            try:
                num = float(val)
            except (TypeError, ValueError):
                skipped += 1
                continue
            try:
                converted = convert(num, category, source_unit, energy_content)
            except (ValueError, KeyError) as e:
                sys.exit(f"sheet '{group['sheet']}', key '{key}': {e}")

            out_rows.append({
                'factor_category': category,
                'factor_key': str(key).strip(),
                'label': str(raw[li]).strip() if li is not None and li < len(raw) and raw[li] is not None else '',
                'co2e_factor': f"{converted:.6g}",
                'unit': CATEGORY_UNIT[category],
                'factor_source': group_source,
                'data_year': data_year,
                'effective_from': eff_from,
                'effective_to': eff_to,
                'notes': note,
            })

    out = base / 'factors.csv'
    with out.open('w', newline='', encoding='utf-8') as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(out_rows)

    print(f"wrote {out} — {len(out_rows)} rows ({skipped} non-numeric or blank rows skipped)")
    print(f"\nNow: ./verify.py {out}")
    print("Then spot-check ten rows by hand against the source document, across\n"
          "different tables. Ten catches both transcription and unit errors; one does not.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)

    i = sub.add_parser('inspect', help='print the sheets, headers and first rows of a workbook')
    i.add_argument('path')
    i.add_argument('--rows', type=int, default=8)
    i.add_argument('--cols', type=int, default=10)
    i.set_defaults(func=cmd_inspect)

    b = sub.add_parser('build', help='emit factors.csv from a mapping.json')
    b.add_argument('mapping')
    b.set_defaults(func=cmd_build)

    args = ap.parse_args()
    return args.func(args)


if __name__ == '__main__':
    sys.exit(main())
